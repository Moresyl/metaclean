"""Verify OOXML property cleanup and document content with independent libraries."""
import argparse
import hashlib
import io
import json
import os
from pathlib import Path
import platform
import re
import shutil
import subprocess
from urllib.parse import urljoin, urlsplit
from xml.etree import ElementTree as ET
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx import __version__ as docx_version
from openpyxl import Workbook, load_workbook
from openpyxl import __version__ as openpyxl_version
from pptx import Presentation
from pptx import __version__ as pptx_version
from pptx.util import Inches

ROOT = Path(__file__).resolve().parent.parent
TIMESTAMP_NS = 1_600_000_000_000_000_000
TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"
RELS_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
ROLES = {
    "core": ("application/vnd.openxmlformats-package.core-properties+xml",
             "http://schemas.openxmlformats.org/package/2006/relationships/metadata/core-properties",
             "http://schemas.openxmlformats.org/package/2006/metadata/core-properties", "coreProperties"),
    "extended": ("application/vnd.openxmlformats-officedocument.extended-properties+xml",
                 "http://schemas.openxmlformats.org/officeDocument/2006/relationships/extended-properties",
                 "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties", "Properties"),
    "custom": ("application/vnd.openxmlformats-officedocument.custom-properties+xml",
               "http://schemas.openxmlformats.org/officeDocument/2006/relationships/custom-properties",
               "http://schemas.openxmlformats.org/officeDocument/2006/custom-properties", "Properties"),
}
ALIASES = {
    "fragment": "properties/private.props#creator",
    "query": "properties/private.props?version=1",
    "query-fragment": "properties/private.props?version=1#creator",
    "unreserved": "properties/%70rivate.props",
    "parent": "../properties/private.props",
    "self-fragment": "properties/private.props",
}


def digest(data):
    return hashlib.sha256(data).hexdigest()


def write_package(path, entries):
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for name, payload in entries.items():
            archive.writestr(name, payload)


def base_package(extension):
    stream = io.BytesIO()
    if extension == "docx":
        document = Document()
        document.add_paragraph("Visible creator and Company text")
        document.add_table(rows=1, cols=1).cell(0, 0).text = "Visible table"
        document.save(stream)
    elif extension == "xlsx":
        workbook = Workbook()
        workbook.active.title = "Visible creator"
        for cell, value in {"A1": 12, "B1": 30, "C1": "=SUM(A1:B1)"}.items():
            workbook.active[cell] = value
        workbook.save(stream)
    else:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        slide.shapes.add_textbox(Inches(1), Inches(1), Inches(7), Inches(1)).text = \
            "Visible creator and Company text"
        presentation.save(stream)
    with ZipFile(stream) as archive:
        return {name: archive.read(name) for name in archive.namelist()}


def make_fixture(path, extension, role, variant):
    entries = base_package(extension)
    content_type, relation, namespace, root_name = ROLES[role]
    if variant == "strict":
        namespace = f"http://purl.oclc.org/ooxml/officeDocument/{role}Properties"
        relation = f"http://purl.oclc.org/ooxml/officeDocument/relationships/{role}Properties"
    location = "properties/private.props"
    if variant == "escaped":
        location = "properties/Private%20Author.props"
    elif variant == "collision":
        location = "customXml/private.props"
    types = ET.fromstring(entries["[Content_Types].xml"])
    old_names = {node.attrib["PartName"].lstrip("/") for node in types
                 if node.attrib.get("ContentType") == content_type}
    for node in list(types):
        if node.attrib.get("ContentType") == content_type:
            types.remove(node)
    for name in old_names:
        entries.pop(name, None)
    ET.SubElement(types, f"{{{TYPES_NS}}}Override", PartName=f"/{location}", ContentType=content_type)
    relationships = ET.fromstring(entries["_rels/.rels"])
    for node in list(relationships):
        if node.attrib.get("Target", "").lstrip("/") in old_names:
            relationships.remove(node)
    ET.SubElement(relationships, f"{{{RELS_NS}}}Relationship", Id="privacyFixture",
                  Type=relation, Target=ALIASES.get(variant, location))
    if variant == "self-fragment":
        child = ET.fromstring(entries["word/_rels/document.xml.rels"])
        ET.SubElement(child, f"{{{RELS_NS}}}Relationship", Id="bookmarkFixture",
                      Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                      Target="#bookmark")
        entries["word/_rels/document.xml.rels"] = ET.tostring(child, encoding="utf-8", xml_declaration=True)
    properties = ET.Element(f"{{{namespace}}}{root_name}")
    if role == "core":
        ET.SubElement(properties, "{http://purl.org/dc/elements/1.1/}creator").text = "Synthetic private author"
    elif role == "extended":
        for tag, value in {"Company": "Synthetic private Company", "Manager": "Synthetic private manager",
                           "Pages": "42"}.items():
            ET.SubElement(properties, f"{{{namespace}}}{tag}").text = value
    else:
        prop = ET.SubElement(properties, f"{{{namespace}}}property",
                             fmtid="{D5CDD505-2E9C-101B-9397-08002B2CF9AE}", pid="2", name="ProjectSecret")
        ET.SubElement(prop, "{http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes}lpwstr").text = \
            "Synthetic private author"
    entries[location] = ET.tostring(properties, encoding="utf-8", xml_declaration=True)
    entries["[Content_Types].xml"] = ET.tostring(types, encoding="utf-8", xml_declaration=True)
    entries["_rels/.rels"] = ET.tostring(relationships, encoding="utf-8", xml_declaration=True)
    write_package(path, entries)
    property_names = {node.attrib["PartName"].lstrip("/") for node in types
                      if node.attrib.get("ContentType") in {values[0] for values in ROLES.values()}}
    return property_names, location, namespace, root_name


def resolve_target(relationship_name, target):
    """Resolve a relationship with stdlib URI operations, independent of the cleaner."""
    source = "" if relationship_name == "_rels/.rels" else \
        relationship_name.replace("/_rels/", "/").removeprefix("_rels/").removesuffix(".rels")
    target = re.sub(r"%([0-9a-fA-F]{2})", lambda match: chr(int(match[1], 16))
                    if chr(int(match[1], 16)) in "ABCDEFGHIJKLMNOPQRSTUVWXYZabcdefghijklmnopqrstuvwxyz0123456789-._~"
                    else match[0], target)
    resolved = urlsplit(urljoin("https://package.test/" + source, target))
    assert resolved.scheme == "https" and resolved.netloc == "package.test", "Unexpected external target"
    return resolved.path.lstrip("/")


def relationship_view(path):
    """Create a library-only view; actual package relationship bytes stay unchanged."""
    view = io.BytesIO()
    changes = []
    with ZipFile(path) as archive, ZipFile(view, "w", compression=ZIP_DEFLATED) as adapted:
        names = {name.lower(): name for name in archive.namelist()}
        assert len(names) == len(archive.namelist()), "Ambiguous part identities"
        for name in archive.namelist():
            payload = archive.read(name)
            if name.endswith(".rels"):
                root = ET.fromstring(payload)
                for node in root:
                    if node.attrib.get("TargetMode") == "External":
                        continue
                    target = node.attrib["Target"]
                    resolved = resolve_target(name, target)
                    assert resolved.lower() in names, f"Unresolved target in {name}"
                    canonical = "/" + names[resolved.lower()]
                    if canonical != target:
                        node.set("Target", canonical)
                        changes.append({"part": name, "id": node.attrib["Id"], "target": target,
                                        "viewTarget": canonical})
                payload = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            adapted.writestr(name, payload)
    view.seek(0)
    return view, changes


def visible_content(stream, extension):
    reopened = io.BytesIO()
    if extension == "docx":
        document = Document(stream)
        before = ([paragraph.text for paragraph in document.paragraphs],
                  [[cell.text for cell in row.cells] for row in document.tables[0].rows])
        assert before == (["Visible creator and Company text"], [["Visible table"]])
        document.save(reopened)
        document = Document(reopened)
        assert ([paragraph.text for paragraph in document.paragraphs],
                [[cell.text for cell in row.cells] for row in document.tables[0].rows]) == before
    elif extension == "xlsx":
        workbook = load_workbook(stream)
        before = (workbook.active.title, [workbook.active[cell].value for cell in ["A1", "B1", "C1"]])
        assert before == ("Visible creator", [12, 30, "=SUM(A1:B1)"])
        workbook.save(reopened)
        workbook = load_workbook(reopened)
        assert (workbook.active.title, [workbook.active[cell].value for cell in ["A1", "B1", "C1"]]) == before
    else:
        presentation = Presentation(stream)
        before = [slide.shapes[0].text for slide in presentation.slides]
        assert before == ["Visible creator and Company text"]
        presentation.save(reopened)
        assert [slide.shapes[0].text for slide in Presentation(reopened).slides] == before
    return before


def check_libraries(original, output, extension, variant):
    direct, errors = [], []
    for path in (original, output):
        try:
            direct.append(visible_content(io.BytesIO(path.read_bytes()), extension))
            errors.append(None)
        except KeyError as error:
            # Only declared URI cases can use an adapter, and both source and
            # output must exhibit the same pre-existing direct-reader limitation.
            if variant not in ALIASES:
                raise
            direct.append(None)
            errors.append(str(error))
    if not any(errors):
        assert direct[0] == direct[1], "Visible document content changed"
        return {"mode": "direct", "openSaveReopen": True}
    assert errors[0] and errors[0] == errors[1], "Library compatibility regressed after cleanup"
    views = [relationship_view(path) for path in (original, output)]
    assert visible_content(views[0][0], extension) == visible_content(views[1][0], extension)
    return {"mode": "relationship-view", "openSaveReopen": True,
            "directReadErrors": {"source": errors[0], "output": errors[1]},
            "viewChanges": {"source": views[0][1], "output": views[1][1]}}


def check_package(original, output, property_names, selected, namespace, root_name):
    with ZipFile(original) as before, ZipFile(output) as after:
        assert before.testzip() is None and after.testzip() is None, "ZIP integrity failed"
        assert len(before.namelist()) == len(set(before.namelist())), "Duplicate source ZIP items"
        assert len(after.namelist()) == len(set(after.namelist())), "Duplicate output ZIP items"
        removed = {name for name in before.namelist()
                   if name.lower().startswith("customxml/") and name not in property_names}
        assert set(after.namelist()) == set(before.namelist()) - removed, "Unexpected package part changes"
        removed_lower = {name.lower() for name in removed}
        for name in after.namelist():
            payload = after.read(name)
            if name in property_names:
                root = ET.fromstring(payload)
                assert len(root) == 0 and not root.attrib and not (root.text or "").strip(), "Private properties remain"
                assert not list(ET.iterparse(io.BytesIO(payload), events=("comment", "pi"))), "Private property markup remains"
                assert root.tag == ET.fromstring(before.read(name)).tag, "Property root identity changed"
                if name == selected:
                    assert root.tag == f"{{{namespace}}}{root_name}", "Property root identity changed"
            elif name == "[Content_Types].xml" or name.endswith(".rels"):
                expected = ET.fromstring(before.read(name))
                changed = False
                for node in list(expected):
                    if name == "[Content_Types].xml":
                        target = node.attrib.get("PartName", "").lstrip("/")
                    elif node.attrib.get("TargetMode") == "External":
                        continue
                    else:
                        target = resolve_target(name, node.attrib["Target"])
                    if target.lower() in removed_lower:
                        expected.remove(node)
                        changed = True
                if changed:
                    actual = ET.fromstring(payload)
                    assert actual.tag == expected.tag and actual.attrib == expected.attrib
                    assert [ET.tostring(node) for node in actual] == [ET.tostring(node) for node in expected], \
                        f"Relationship or manifest records changed: {name}"
                else:
                    assert payload == before.read(name), f"Unrelated reference bytes changed: {name}"
            else:
                assert payload == before.read(name), f"Visible or unrelated payload changed: {name}"
        return sorted(removed)


def run_native(arguments, log, environment=None):
    result = subprocess.run(arguments, cwd=ROOT, capture_output=True, text=True,
                            encoding="utf-8", errors="replace", timeout=600 if environment is None else 120,
                            env=environment)
    log.write_text(result.stdout + result.stderr, encoding="utf-8")
    assert result.returncode == 0, result.stdout + result.stderr
    return result


def verify_case(directory, extension, role, variant, mode, identities):
    case = directory / mode
    case.mkdir()
    source = case / f"sample.{extension}"
    shutil.copyfile(directory / f"original.{extension}", source)
    os.utime(source, ns=(TIMESTAMP_NS, TIMESTAMP_NS))
    before = source.read_bytes()
    command = ["cargo", "test", "--locked", "--manifest-path", "src-tauri/Cargo.toml", "--lib",
               "cleans_external_ooxml_sample_with_verified_output", "--", "--ignored"]
    result = run_native(command, case / "native.log",
                        {**os.environ, "METACLEAN_OOXML_SAMPLE_PATH": str(source), "METACLEAN_OOXML_OUTPUT_MODE": mode})
    assert "1 passed; 0 failed" in result.stdout, "Native fixture test did not execute"
    output = source if mode == "replace" else case / f"sample.cleaned.{extension}"
    original = case / f"sample.{extension}.bak" if mode == "replace" else source
    assert original.read_bytes() == before, "Source or replacement backup changed"
    assert original.stat().st_mtime_ns == TIMESTAMP_NS and output.stat().st_mtime_ns == TIMESTAMP_NS
    removed = check_package(original, output, *identities)
    libraries = check_libraries(original, output, extension, variant)
    return {"extension": extension, "role": role, "variant": variant, "mode": mode,
            "sourceSha256": digest(before), "outputSha256": digest(output.read_bytes()),
            "sourceOrBackupExact": True, "filesystemMtimePreserved": True,
            "propertiesCleared": True, "survivingPayloadsExact": True, "referencesPrunedExactly": True,
            "removedCustomXml": removed, "library": libraries}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    (output / "tools.json").write_text(json.dumps({
        "python": platform.python_version(), "python-docx": docx_version,
        "openpyxl": openpyxl_version, "python-pptx": pptx_version,
    }, indent=2), encoding="utf-8")
    scenarios = [(extension, role, "relocated") for extension in ("docx", "xlsx", "pptx") for role in ROLES]
    scenarios += [("docx", role, "strict") for role in ("extended", "custom")]
    scenarios += [("docx", "core", "escaped")]
    scenarios += [("docx", role, "collision") for role in ROLES]
    scenarios += [("docx", "core", alias) for alias in ALIASES]
    results = []
    try:
        run_native(["cargo", "test", "--locked", "--manifest-path", "src-tauri/Cargo.toml", "--lib", "--no-run"],
                   output / "compile.log")
        for extension, role, variant in scenarios:
            directory = output / f"{extension}-{role}-{variant}"
            directory.mkdir()
            identities = make_fixture(directory / f"original.{extension}", extension, role, variant)
            for mode in ("copy", "replace"):
                result = verify_case(directory, extension, role, variant, mode, identities)
                results.append(result)
                print(f"Verified {extension}/{role}/{variant}/{mode}: properties cleared, payloads exact, "
                      f"library={result['library']['mode']}", flush=True)
        assert len(results) == 42, "Incomplete qualification"
    finally:
        (output / "results.json").write_text(json.dumps(results, indent=2), encoding="utf-8")


if __name__ == "__main__":
    main()
